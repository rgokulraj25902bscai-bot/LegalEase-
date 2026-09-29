from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF
from core.text_utils import pdf_safe_text, sanitize_text

def to_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")

def to_docx(text: str, document_type: str = "Legal Document") -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(document_type)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    document.add_paragraph()
    for line in sanitize_text(text).split("\n"):
        stripped = line.strip()
        if not stripped:
            document.add_paragraph()
            continue
        if stripped.startswith("### "):
            paragraph = document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            paragraph = document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            paragraph = document.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            paragraph = document.add_paragraph(stripped)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Generated with LegalEase")
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()

class LegalEasePDF(FPDF):
    def __init__(self, document_type: str):
        super().__init__()
        self.document_type = document_type
        self.set_auto_page_break(auto=True, margin=18)
    def header(self):
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 8, pdf_safe_text(self.document_type), align="C")
        self.ln(12)
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 8, f"Page {self.page_no()}", align="C")

def to_pdf(text: str, document_type: str = "Legal Document") -> bytes:
    """Export the document to a PDF using a stable, explicit text width."""
    pdf = LegalEasePDF(document_type)
    pdf.set_margins(18, 18, 18)
    pdf.add_page()

    # Explicitly reset X before every multi_cell. This avoids fpdf2's
    # "Not enough horizontal space to render a single character" error
    # when the cursor is left near the right margin by a previous cell.
    text_width = pdf.w - pdf.l_margin - pdf.r_margin

    for line in sanitize_text(text).split("\n"):
        stripped = line.strip()
        if not stripped:
            pdf.set_x(pdf.l_margin)
            pdf.ln(5)
            continue

        if stripped.startswith("### "):
            content, size, leading = stripped[4:], 12, 7
            bold = True
        elif stripped.startswith("## "):
            content, size, leading = stripped[3:], 13, 8
            bold = True
        elif stripped.startswith("# "):
            content, size, leading = stripped[2:], 15, 9
            bold = True
        elif stripped.startswith("- "):
            content, size, leading = "- " + stripped[2:], 11, 6
            bold = False
        else:
            content, size, leading = stripped, 11, 6
            bold = False

        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B" if bold else "", size)
        pdf.multi_cell(text_width, leading, pdf_safe_text(content), new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())
